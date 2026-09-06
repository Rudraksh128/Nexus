import csv
import io
import json
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.ingestion.models import ParsedPage


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".docx",
    ".csv",
    ".json",
    ".xlsx",
}


def get_extension(filename: str) -> str:
    """
    Return a normalized lowercase file extension.
    """
    return Path(filename).suffix.lower()


def is_supported_file(filename: str) -> bool:
    """
    Check whether NEXUS supports the supplied file type.
    """
    return get_extension(filename) in SUPPORTED_EXTENSIONS


def _parse_pdf(file_path: Path) -> list[ParsedPage]:
    reader = PdfReader(str(file_path))

    pages: list[ParsedPage] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = (page.extract_text() or "").strip()

        if text:
            pages.append(
                ParsedPage(
                    page_number=page_number,
                    text=text,
                )
            )

    return pages


def _parse_text(file_path: Path) -> list[ParsedPage]:
    text = file_path.read_text(
        encoding="utf-8",
        errors="replace",
    ).strip()

    if not text:
        return []

    return [
        ParsedPage(
            page_number=1,
            text=text,
        )
    ]


def _parse_docx(file_path: Path) -> list[ParsedPage]:
    document = DocxDocument(str(file_path))

    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            blocks.append(text)

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        rows: list[str] = [
            f"TABLE {table_index}"
        ]

        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            rows.append(
                " | ".join(cells)
            )

        blocks.append(
            "\n".join(rows)
        )

    text = "\n\n".join(blocks).strip()

    if not text:
        return []

    return [
        ParsedPage(
            page_number=1,
            text=text,
        )
    ]


def _parse_csv(file_path: Path) -> list[ParsedPage]:
    rows: list[str] = []

    with file_path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline="",
    ) as file:
        reader = csv.reader(file)

        for row in reader:
            rows.append(
                " | ".join(
                    cell.strip()
                    for cell in row
                )
            )

    text = "\n".join(rows).strip()

    if not text:
        return []

    return [
        ParsedPage(
            page_number=1,
            text=text,
        )
    ]


def _parse_json(file_path: Path) -> list[ParsedPage]:
    raw_text = file_path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    data = json.loads(raw_text)

    formatted = json.dumps(
        data,
        indent=2,
        ensure_ascii=False,
    )

    if not formatted.strip():
        return []

    return [
        ParsedPage(
            page_number=1,
            text=formatted,
        )
    ]


def _parse_xlsx(file_path: Path) -> list[ParsedPage]:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    pages: list[ParsedPage] = []

    page_number = 0

    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = [
                f"SHEET: {worksheet.title}"
            ]

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = []

                for value in row:
                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value))

                if any(value.strip() for value in values):
                    rows.append(
                        " | ".join(values)
                    )

            text = "\n".join(rows).strip()

            if text:
                page_number += 1

                pages.append(
                    ParsedPage(
                        page_number=page_number,
                        text=text,
                    )
                )

    finally:
        workbook.close()

    return pages


def parse_document(
    file_path: str | Path,
    filename: str,
) -> list[ParsedPage]:
    """
    Parse a supported document into a common page representation.
    """

    path = Path(file_path)
    extension = get_extension(filename)

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":
        return _parse_pdf(path)

    if extension in {".txt", ".md"}:
        return _parse_text(path)

    if extension == ".docx":
        return _parse_docx(path)

    if extension == ".csv":
        return _parse_csv(path)

    if extension == ".json":
        return _parse_json(path)

    if extension == ".xlsx":
        return _parse_xlsx(path)

    raise ValueError(
        f"No parser registered for {extension}"
    )